#!/usr/bin/env python3
"""Generate synthetic exports, independently read text, and rasterize PDF pages.

This is not visual DOCX verification: only Word/LibreOffice can paginate DOCX.
No network, databases, model calls, or user meeting data are used.
"""

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pypdfium2 as pdfium
from docx import Document
from pypdf import PdfReader

from sozvon.export.document import Block, ExportDocument
from sozvon.export.docx import render_docx
from sozvon.export.pdf import render_pdf

FONTS = ROOT / "sozvon/web/static/fonts"


def samples():
    short = ExportDocument("Встреча — проект Ёж", (
        Block("meta", "СИНТЕТИЧЕСКИЙ ПРИМЕР · Не данные реальной встречи"),
        Block("h2", "Решения"),
        Block("p", "Анна & Илья: проверка <условия> выполнена."),
        Block("h2", "Задачи"),
        Block("p", "Обновить инструкцию и проверить экспорт кириллицы: Ёж, съёмка, объём."),
        Block("meta", "Исполнитель: Не назначен · Срок: Не указан"),
        Block("quote", "[Абзац 2] Срок пока не обещаю."),
        Block("h2", "Буквальный текст, не HTML"),
        Block("p", '<img src="https://example.invalid/test.png"/> & <script>пример</script>'),
        Block("h2", "Расшифровка"),
        Block("meta", "[00:01:05] · Анна"),
        Block("p", "Первая строка.\nВторая строка, явный перенос сохранён."),
    ))
    blocks = list(short.blocks)
    blocks.append(Block("h2", "Многостраничная проверка"))
    for index in range(110):
        blocks.append(Block("p", f"Пункт {index}. " + (
            "Синтетический абзац проверяет переносы, кириллицу и сохранность текста. " * 12)))
        if index % 15 == 0:
            blocks.append(Block("quote", f"[Абзац {index + 1}] Источник синтетического пункта."))
    blocks.extend((
        Block("h2", "Длинное слово"),
        Block("p", "Сверхдлинноеслово" * 60),
        Block("p", "ПОСЛЕДНИЙ АБЗАЦ — проверка завершена."),
    ))
    return {"short": short, "long": ExportDocument("Большой синтетический созвон", tuple(blocks))}


def verify(output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    manifest = {"synthetic": True, "docx_visual_review": "not_performed", "files": []}
    for name, document in samples().items():
        expected = [document.title, *(block.text for block in document.blocks)]
        for fmt in ("docx", "pdf"):
            target = output / f"{name}.{fmt}"
            record = {"filename": target.name, "pages": None, "images": []}
            if fmt == "docx":
                render_docx(document, target)
                readback = Document(target)
                paragraphs = [paragraph.text for paragraph in readback.paragraphs]
                if paragraphs != expected:
                    raise AssertionError(f"DOCX paragraph mismatch: {target.name}")
                text = "\n".join(paragraphs)
            else:
                render_pdf(document, target, FONTS)
                reader = PdfReader(target)
                record["pages"] = len(reader.pages)
                bodies = []
                for page_number, page in enumerate(reader.pages, 1):
                    footer, body = (page.extract_text() or "").split("\n", 1)
                    if footer != str(page_number):
                        raise AssertionError(f"Missing page number: {target.name}:{page_number}")
                    bodies.append(body)
                    fonts = [font.get_object()
                             for font in page["/Resources"]["/Font"].values()]
                    if not any("/FontDescriptor" in font and
                               "/FontFile2" in font["/FontDescriptor"] for font in fonts):
                        raise AssertionError(f"Missing embedded font: {target.name}:{page_number}")
                text = "".join(bodies)
                with pdfium.PdfDocument(target) as pdf:
                    if len(pdf) != len(reader.pages):
                        raise AssertionError("Independent readers disagree on PDF page count")
                    for page_index in range(len(pdf)):
                        image_name = f"{name}-page-{page_index + 1:03d}.png"
                        page = pdf[page_index]
                        try:
                            bitmap = page.render(scale=1.5)
                            try:
                                bitmap.to_pil().save(output / image_name)
                            finally:
                                bitmap.close()
                        finally:
                            page.close()
                        record["images"].append(image_name)
            compact = "".join(text.split())
            missing = [value for value in expected if "".join(value.split()) not in compact]
            record["bytes"] = target.stat().st_size
            record["missing_strings"] = len(missing)
            manifest["files"].append(record)
            if missing:
                raise AssertionError(f"Missing text in {target.name}: {len(missing)} blocks")
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                                         encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/export-qa")
    args = parser.parse_args()
    manifest = verify(args.output.resolve())
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
