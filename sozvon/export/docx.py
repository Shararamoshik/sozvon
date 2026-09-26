"""Direct OOXML renderer; all content is literal text, never XML."""

from pathlib import Path
from typing import BinaryIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from sozvon.export.document import ExportDocument


def render_docx(document: ExportDocument, target: Path | BinaryIO) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + kind), "Arial")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    doc.core_properties.title = document.title
    doc.core_properties.author = "Созвон"
    doc.add_heading(document.title, level=0)
    for block in document.blocks:
        if block.kind == "h2":
            paragraph = doc.add_heading(block.text, level=1)
            paragraph.paragraph_format.keep_with_next = True
        elif block.kind in {"p", "quote", "meta"}:
            paragraph = doc.add_paragraph(block.text)
            if block.kind == "quote":
                paragraph.paragraph_format.left_indent = Cm(.5)
            if block.kind == "meta":
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        else:
            raise ValueError("Неизвестный блок документа")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(str(target) if isinstance(target, Path) else target)
