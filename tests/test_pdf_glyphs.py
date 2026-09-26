"""Real PDF/font coverage failures and independent document readback."""

import json
from dataclasses import asdict
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from sozvon.export.document import Block, ExportDocument
from sozvon.export.pdf import render_pdf
from sozvon.web.server import create_app

FONTS = Path(__file__).resolve().parents[1] / "sozvon/web/static/fonts"


@pytest.mark.parametrize("kind", ["title", "h2", "p", "meta", "quote"])
@pytest.mark.parametrize("char", ["α", "✅", "\U0001f680", "\u2007", "\u200b", "\u2060", "\u3000"])
def test_missing_glyph_in_every_style_fails_before_writing(tmp_path, kind, char):
    private = f"Частный текст {char} не включать в ошибку"
    title = private if kind == "title" else "Заголовок"
    # Put the failure after valid content to exercise whole-document preflight.
    blocks = (Block("p", "Первый допустимый абзац"),
              Block("p" if kind == "title" else kind, private if kind != "title" else "Текст"))
    document = ExportDocument(title, blocks)
    fresh = tmp_path / "new.pdf"
    existing = tmp_path / "existing.pdf"
    existing.write_bytes(b"original file must remain untouched")
    output = BytesIO(b"original stream must remain untouched")
    original_file, original_stream = existing.read_bytes(), output.getvalue()
    for target in (fresh, existing, output):
        with pytest.raises(ValueError) as failure:
            render_pdf(document, target, FONTS)
        detail = str(failure.value)
        assert f"U+{ord(char):04X}" in detail and "DOCX" in detail
        assert "Частный" not in detail and private not in detail
    assert not fresh.exists()
    assert existing.read_bytes() == original_file
    assert output.getvalue() == original_stream and output.tell() == 0


def test_supported_characters_have_exact_independent_pdf_readback(tmp_path):
    title = "Русский Ёж — отчёт № 2"
    values = ("Решения: «Ёж» & <условие>", "Кириллица ёж / English 123: ± × ÷ =",
              "Стоимость: 50 ₽; 10 €; 5 £; 20 %", "Цитата — ‘слова’ © ® ™…")
    document = ExportDocument(title, tuple(
        Block(kind, value) for kind, value in zip(("h2", "p", "meta", "quote"), values)
    ))
    path = tmp_path / "supported.pdf"
    render_pdf(document, path, FONTS)
    reader = PdfReader(path)
    readback = "\n".join(page.extract_text() for page in reader.pages)
    assert "\x00" not in readback and "\ufffd" not in readback
    for value in (title, *values):
        assert value in readback
    assert reader.metadata is not None and reader.metadata.title == title


def test_pdf_layout_controls_are_not_mistaken_for_missing_glyphs(tmp_path):
    document = ExportDocument("Заголовок\r\nВторая строка", (
        Block("p", "Первая\tстрока\nВторая строка\rТретья строка"),
    ))
    path = tmp_path / "controls.pdf"
    render_pdf(document, path, FONTS)
    readback = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    assert "\x00" not in readback
    assert "Заголовок\nВторая строка" in readback
    assert "Первая строка\nВторая строка\nТретья строка" in readback


@pytest.mark.parametrize("char", ["α", "✅", "\u2007"])
def test_real_pdf_worker_keeps_previous_output_on_missing_glyph(tmp_path, char):
    from sozvon.runtime.host import WorkerHost
    from sozvon.runtime.protocol import WorkerError

    target = tmp_path / "output.pdf"
    render_pdf(ExportDocument("Старый документ", (Block("p", "Исходный результат"),)),
               target, FONTS)
    previous = target.read_bytes()
    invalid = ExportDocument("Новый документ", (Block("quote", f"Частная цитата {char}"),))
    source = tmp_path / "input.json"
    source.write_text(json.dumps(asdict(invalid), ensure_ascii=False), encoding="utf-8")
    snapshot = source.read_bytes()
    with pytest.raises(WorkerError) as failure:
        WorkerHost().run("render_export", {"job_dir": str(tmp_path), "format": "pdf"},
                         timeout=20)
    assert failure.value.code == "operation_error"
    assert f"U+{ord(char):04X}" in str(failure.value) and "DOCX" in str(failure.value)
    assert "Частная цитата" not in str(failure.value)
    assert target.read_bytes() == previous
    assert source.read_bytes() == snapshot
    assert set(tmp_path.iterdir()) == {target, source}
    text = "\n".join(page.extract_text() for page in PdfReader(target).pages)
    assert "Исходный результат" in text and "Новый документ" not in text


def test_pdf_api_rejects_missing_glyph_without_publishing_or_mutating_source(tmp_path):
    text = "Русский Ёж: α + β = λ. Статус ✅"
    app = create_app(tmp_path)
    repo = app.state.service.repo
    mid = repo.create_text("Синтетическая встреча", text)
    source = tmp_path / "original.txt"
    source.write_text(text, encoding="utf-8")
    original = source.read_bytes()
    repo.attach_source(mid, source, "text")
    before = repo.detail(mid)

    with TestClient(app, base_url="http://127.0.0.1") as client:
        login = client.post("/api/session", json={"key": app.state.launch_key},
                            headers={"Origin": "http://127.0.0.1"})
        assert login.status_code == 200
        response = client.get(f"/api/meetings/{mid}/export?format=pdf")
        assert response.status_code == 422
        assert "content-disposition" not in response.headers
        assert response.headers["content-type"].startswith("application/json")
        detail = response.json()["detail"]
        assert "U+03B1" in detail and "DOCX" in detail
        assert "шрифт" in detail.lower()
        assert text not in detail and "Русский Ёж" not in detail
        assert list((tmp_path / "exports" / "tmp").iterdir()) == []
        assert repo.detail(mid) == before
        assert source.read_bytes() == original

        # The suggested alternative must preserve the entire source, not sanitize it.
        alternative = client.get(f"/api/meetings/{mid}/export?format=docx")
        assert alternative.status_code == 200
        readback = [p.text for p in Document(BytesIO(alternative.content)).paragraphs]
        assert text in readback
        assert repo.detail(mid) == before
        assert source.read_bytes() == original
        assert list((tmp_path / "exports" / "tmp").iterdir()) == []
