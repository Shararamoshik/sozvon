"""Export worker boundary, using synthetic snapshots and real child processes."""

import json
import os
from pathlib import Path

import pytest
from docx import Document

from sozvon.runtime.host import WorkerHost


def write_snapshot(folder, value=None):
    value = value or {"title": "Синтетический созвон Ёж", "blocks": [
        {"kind": "p", "text": "Завершение синтетического документа & <текст>"},
    ]}
    source = folder / "input.json"
    source.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    return source


def test_real_worker_docx_returns_only_name_and_size(tmp_path):
    source = write_snapshot(tmp_path)
    original = source.read_bytes()
    events = []
    result = WorkerHost().run("render_export", {"job_dir": str(tmp_path), "format": "docx"},
                              timeout=20, on_event=events.append)
    target = tmp_path / "output.docx"
    assert result == {"filename": "output.docx", "size": target.stat().st_size}
    assert target.read_bytes().startswith(b"PK")
    doc = Document(target)
    assert doc.paragraphs[0].text == "Синтетический созвон Ёж"
    assert doc.paragraphs[-1].text == "Завершение синтетического документа & <текст>"
    assert source.read_bytes() == original
    assert not (tmp_path / "output.part").exists()
    assert next(event["pid"] for event in events if event["type"] == "ready") != os.getpid()
    assert any(event["type"] == "progress" for event in events)
    assert Path(result["filename"]).name == result["filename"]


def test_real_worker_pdf_preserves_literal_text(tmp_path):
    from pypdf import PdfReader

    write_snapshot(tmp_path)
    result = WorkerHost().run("render_export", {"job_dir": str(tmp_path), "format": "pdf"},
                              timeout=20)
    target = tmp_path / "output.pdf"
    assert result == {"filename": "output.pdf", "size": target.stat().st_size}
    text = "\n".join(page.extract_text() for page in PdfReader(target).pages)
    assert "Синтетический созвон Ёж" in text
    assert "Завершение синтетического документа & <текст>" in text


@pytest.mark.parametrize("extra", [
    {"format": "html"}, {"format": "../../escape"}, {"format": []},
    {"module": "os"}, {"target": "outside.pdf"}, {"url": "https://example.invalid"},
])
def test_worker_rejects_untrusted_payload_fields(tmp_path, extra):
    from threading import Event

    from sozvon.export.worker import render_export

    source = write_snapshot(tmp_path)
    payload = {"job_dir": str(tmp_path), "format": "docx", **extra}
    with pytest.raises(ValueError):
        render_export(payload, Event(), lambda event: None)
    assert set(tmp_path.iterdir()) == {source}


@pytest.mark.parametrize("value", [
    [], {}, {"title": "Документ", "blocks": [], "extra": 1},
    {"title": "Документ", "blocks": {"kind": "p", "text": "Текст"}},
    {"title": "Документ", "blocks": [{"kind": "p", "text": "Текст", "extra": 1}]},
    {"title": "Документ", "blocks": [7]},
])
def test_worker_validates_snapshot_shape(tmp_path, value):
    from threading import Event

    from sozvon.export.worker import render_export

    (tmp_path / "input.json").write_text(json.dumps(value), encoding="utf-8")
    with pytest.raises(ValueError, match="снимок|блок"):
        render_export({"job_dir": str(tmp_path), "format": "pdf"}, Event(), lambda event: None)
    assert not (tmp_path / "output.pdf").exists()


def test_worker_rejects_input_over_eight_mib_before_parse(tmp_path):
    from threading import Event

    from sozvon.export.worker import render_export

    source = write_snapshot(tmp_path)
    with source.open("ab") as stream:
        stream.write(b" " * (8 * 1024 * 1024))
    with pytest.raises(ValueError, match="больш"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)
    assert not (tmp_path / "output.docx").exists()


@pytest.mark.parametrize("data", [
    b'{"title":"One","title":"Two","blocks":[{"kind":"p","text":"Text"}]}',
    b'{"title":NaN,"blocks":[{"kind":"p","text":"Text"}]}',
    b'[' * 2000 + b']' * 2000,
    b'\xff',
])
def test_worker_rejects_ambiguous_json_safely(tmp_path, data):
    from threading import Event

    from sozvon.export.worker import render_export

    (tmp_path / "input.json").write_bytes(data)
    with pytest.raises(ValueError, match="Некорректный экспортный снимок"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)


@pytest.mark.parametrize("stage", ["before", "progress", "after_render"])
def test_cancelled_worker_never_publishes_partial_document(tmp_path, monkeypatch, stage):
    from threading import Event

    from sozvon.export import docx
    from sozvon.export.worker import render_export

    write_snapshot(tmp_path)
    target = tmp_path / "output.docx"
    target.write_bytes(b"previous export must survive cancellation")
    stop = Event()
    original_render = docx.render_docx

    def render(document, output):
        original_render(document, output)
        if stage == "after_render":
            stop.set()

    monkeypatch.setattr(docx, "render_docx", render)
    if stage == "before":
        stop.set()

    def emit(event):
        if stage == "progress":
            stop.set()

    with pytest.raises(RuntimeError, match="отмен"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, stop, emit)
    assert target.read_bytes() == b"previous export must survive cancellation"
    assert not (tmp_path / "output.part").exists()


def test_oversized_output_is_not_published(tmp_path, monkeypatch):
    from threading import Event

    from sozvon.export import docx
    from sozvon.export.worker import render_export

    write_snapshot(tmp_path)
    target = tmp_path / "output.docx"
    target.write_bytes(b"previous valid result")

    def oversized_renderer(document, path):
        # Fault injection at the renderer boundary, not a claimed valid DOCX.
        path.truncate(32 * 1024 * 1024 + 1)

    monkeypatch.setattr(docx, "render_docx", oversized_renderer)
    with pytest.raises(ValueError, match="больш"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)
    assert target.read_bytes() == b"previous valid result"
    assert not (tmp_path / "output.part").exists()


@pytest.mark.parametrize("name", ["input.json", "output.part", "output.docx"])
def test_worker_refuses_symlinks_without_touching_outside_files(tmp_path, name):
    from threading import Event

    from sozvon.export.worker import render_export

    folder = tmp_path / "job"
    folder.mkdir()
    source = write_snapshot(folder)
    outside = tmp_path / "outside"
    outside.write_bytes(source.read_bytes())
    original = outside.read_bytes()
    link = folder / name
    if link.exists():
        link.unlink()
    link.symlink_to(outside)
    with pytest.raises(ValueError, match="файл|ссылк"):
        render_export({"job_dir": str(folder), "format": "docx"}, Event(), lambda event: None)
    assert outside.read_bytes() == original
    assert link.is_symlink()


def test_worker_rejects_linked_job_directory(tmp_path):
    from threading import Event

    from sozvon.export.worker import render_export

    real = tmp_path / "real"
    real.mkdir()
    write_snapshot(real)
    linked = tmp_path / "linked"
    linked.symlink_to(real, target_is_directory=True)
    with pytest.raises(ValueError, match="папк|ссылк"):
        render_export({"job_dir": str(linked), "format": "docx"}, Event(), lambda event: None)
    assert not (real / "output.docx").exists()


def test_real_worker_timeout_reaps_child_without_publishing(tmp_path, monkeypatch):
    from sozvon.runtime import host
    from sozvon.runtime.protocol import WorkerError

    write_snapshot(tmp_path, {"title": "Синтетический стресс-тест", "blocks": [
        {"kind": "p", "text": "Длинный текст для отмены. " * 50}
        for _ in range(2000)
    ]})
    real_popen = host.subprocess.Popen
    children = []

    def observed_popen(*args, **kwargs):
        child = real_popen(*args, **kwargs)
        children.append(child)
        return child

    monkeypatch.setattr(host.subprocess, "Popen", observed_popen)
    with pytest.raises(WorkerError) as exc:
        WorkerHost(stop_grace=.05).run(
            "render_export", {"job_dir": str(tmp_path), "format": "pdf"}, timeout=.03)
    assert exc.value.code == "timeout"
    assert len(children) == 1 and children[0].poll() is not None
    assert not (tmp_path / "output.pdf").exists()


def test_worker_rechecks_output_before_opening_after_progress(tmp_path):
    from threading import Event

    from sozvon.export.worker import render_export

    write_snapshot(tmp_path)
    outside = tmp_path / "untouched"
    outside.write_bytes(b"outside must not be overwritten")
    partial = tmp_path / "output.part"

    def emit(event):
        partial.symlink_to(outside)

    with pytest.raises(ValueError, match="файл|ссылк|занят"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), emit)
    assert outside.read_bytes() == b"outside must not be overwritten"
    assert partial.is_symlink()


@pytest.mark.parametrize("path", ["relative", "https://example.invalid/job", "//server/share/job"])
def test_worker_rejects_nonlocal_job_paths(path):
    from threading import Event

    from sozvon.export.worker import render_export

    with pytest.raises(ValueError, match="локальный"):
        render_export({"job_dir": path, "format": "docx"}, Event(), lambda event: None)


def test_worker_rejects_source_replacement_between_check_and_open(tmp_path, monkeypatch):
    from threading import Event

    from sozvon.export.worker import render_export

    source = write_snapshot(tmp_path)
    outside = tmp_path / "private.json"
    outside.write_bytes(source.read_bytes())
    original_open = Path.open
    triggered = False

    def racing_open(path, *args, **kwargs):
        nonlocal triggered
        if path == source and not triggered:
            triggered = True
            path.unlink()
            path.symlink_to(outside)
        return original_open(path, *args, **kwargs)

    monkeypatch.setattr(Path, "open", racing_open)
    with pytest.raises(ValueError, match="файл|ссылк|измен"):
        render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)
    assert triggered
    assert not (tmp_path / "output.docx").exists()


def test_worker_never_publishes_replaced_temporary_file(tmp_path, monkeypatch):
    from threading import Event

    from sozvon.export import docx
    from sozvon.export.worker import render_export

    write_snapshot(tmp_path)
    target = tmp_path / "output.docx"
    target.write_bytes(b"previous export")
    outside = tmp_path / "outside"
    outside.write_bytes(b"outside export")
    original_render = docx.render_docx

    def replace_after_render(document, output):
        original_render(document, output)
        path = Path(output.name)
        path.unlink()
        path.symlink_to(outside)

    monkeypatch.setattr(docx, "render_docx", replace_after_render)
    import os
    if os.name == "nt":
        # Windows denies unlink of the open file before the injected substitution.
        with pytest.raises(PermissionError) as caught:
            render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)
        assert caught.value.winerror == 32
    else:
        with pytest.raises(ValueError, match="измен|файл"):
            render_export({"job_dir": str(tmp_path), "format": "docx"}, Event(), lambda event: None)
    assert target.read_bytes() == b"previous export"
    assert not target.is_symlink()
    assert outside.read_bytes() == b"outside export"
