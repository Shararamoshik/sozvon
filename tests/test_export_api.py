"""Router tests compose the new router with the existing session middleware."""

import importlib
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from sozvon.web.server import create_app


def app_with_router(tmp_path, worker_factory=None):
    router = importlib.import_module("sozvon.web.routes.export")
    app = create_app(tmp_path, worker_factory=worker_factory)
    # Parent owns server registration. Compose explicitly until it replaces the legacy route.
    app.router.routes[:] = [r for r in app.router.routes
                           if getattr(r, "path", None) != "/api/meetings/{mid}/export"]
    app.include_router(router.create_router(app.state.service))
    return app


def login(app, client):
    response = client.post("/api/session", json={"key": app.state.launch_key},
                           headers={"Origin": "http://127.0.0.1"})
    assert response.status_code == 200


def test_docx_api_preserves_session_and_safe_download_headers(tmp_path):
    app = app_with_router(tmp_path)
    service = app.state.service
    mid = service.repo.create_text('Русское имя "../evil"', "Синтетический текст Ёж")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        url = f"/api/meetings/{mid}/export?format=docx"
        assert client.get(url).status_code == 401
        login(app, client)
        response = client.get(url)
        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert response.headers["content-disposition"] == (
            f'attachment; filename="sozvon-{mid}-transcript.docx"')
        assert response.headers["cache-control"] == "no-store"
        assert response.headers["x-content-type-options"] == "nosniff"
        assert "Синтетический текст Ёж" in "\n".join(
            p.text for p in Document(BytesIO(response.content)).paragraphs)


@pytest.mark.parametrize("query", ["format=html", "content=all", "include_quotes=false",
                                  "include_transcript=yes", "report_id=invalid"])
def test_invalid_export_query_returns_422_not_conflict(tmp_path, query):
    app = app_with_router(tmp_path)
    mid = app.state.service.repo.create_text("Ёж", "Текст")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        response = client.get(f"/api/meetings/{mid}/export?{query}")
        assert response.status_code == 422
        assert "content-disposition" not in response.headers


def test_snapshot_limit_is_http_413(tmp_path):
    app = app_with_router(tmp_path)
    mid = app.state.service.repo.create_text("Ёж", "Текст")
    with app.state.service.repo.connection() as conn:
        conn.execute("UPDATE meetings SET notes=? WHERE id=?", ("x" * (8 * 1024 * 1024), mid))
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        response = client.get(f"/api/meetings/{mid}/export?format=pdf")
        assert response.status_code == 413
        assert "content-disposition" not in response.headers


@pytest.mark.parametrize("reason", ["disconnect", "task_cancel"])
def test_http_cancellation_reaps_real_worker_before_request_finishes(tmp_path, monkeypatch, reason):
    import asyncio
    import threading

    from fastapi import FastAPI

    from sozvon.runtime import host
    from sozvon.services.application import Application
    from sozvon.web.routes.export import create_router

    rendering = threading.Event()
    children = []
    original = host.subprocess.Popen

    def observe(*args, **kwargs):
        child = original(*args, **kwargs)
        children.append(child)
        return child

    class ObservedWorker(host.WorkerHost):
        def run(self, *args, **kwargs):
            def on_event(event):
                if event["type"] == "progress":
                    rendering.set()
            return super().run(*args, **kwargs, on_event=on_event)

    monkeypatch.setattr(host.subprocess, "Popen", observe)
    service = Application(tmp_path, worker_factory=lambda: ObservedWorker(stop_grace=.05))
    mid = service.repo.create_text("HTTP отмена", "\n".join(["Длинная строка. " * 50] * 500))
    app = FastAPI()
    app.include_router(create_router(service))
    messages = []

    async def scenario():
        incoming = asyncio.Queue()
        await incoming.put({"type": "http.request", "body": b"", "more_body": False})

        async def send(message):
            messages.append(message)

        scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
                 "method": "GET", "scheme": "http", "path": f"/api/meetings/{mid}/export",
                 "raw_path": f"/api/meetings/{mid}/export".encode(), "query_string": b"format=pdf",
                 "root_path": "", "headers": [], "client": ("127.0.0.1", 50000),
                 "server": ("127.0.0.1", 8000)}
        request = asyncio.create_task(app(scope, incoming.get, send))
        try:
            assert await asyncio.to_thread(rendering.wait, 5)
            if reason == "disconnect":
                await incoming.put({"type": "http.disconnect"})
                # asyncio.wait does not cancel on timeout: failure cannot accidentally reap the child.
                done, _ = await asyncio.wait({request}, timeout=3)
                assert request in done, "HTTP disconnect must stop export without task cancellation"
                await request
                assert not any(m.get("status") == 200 for m in messages)
            else:
                request.cancel()
                with pytest.raises(asyncio.CancelledError):
                    await request
            assert len(children) == 1 and children[0].poll() is not None
            assert list((tmp_path / "exports" / "tmp").iterdir()) == []
            assert service._export_lock.acquire(blocking=False)
            service._export_lock.release()
            assert not service._export_stop.is_set()
        finally:
            service._export_stop.set()
            if not request.done():
                await asyncio.wait_for(asyncio.gather(request, return_exceptions=True), timeout=5)

    asyncio.run(scenario())
    service.close()


def test_real_socket_disconnect_through_session_middleware_reaps_worker(tmp_path, monkeypatch):
    import socket
    import threading
    import time

    import httpx
    import uvicorn

    from sozvon.runtime import host

    rendering, finished = threading.Event(), threading.Event()
    children = []
    original = host.subprocess.Popen

    def observe(*args, **kwargs):
        child = original(*args, **kwargs)
        children.append(child)
        return child

    class ObservedWorker(host.WorkerHost):
        def run(self, *args, **kwargs):
            def on_event(event):
                if event["type"] == "progress":
                    rendering.set()
            try:
                return super().run(*args, **kwargs, on_event=on_event)
            finally:
                finished.set()

    monkeypatch.setattr(host.subprocess, "Popen", observe)
    app = app_with_router(tmp_path, lambda: ObservedWorker(stop_grace=.05))
    service = app.state.service
    mid = service.repo.create_text("Сокет отмена", "\n".join(["Длинная строка. " * 50] * 500))
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
    thread.start()
    try:
        deadline = time.monotonic() + 5
        while not server.started and time.monotonic() < deadline:
            time.sleep(.01)
        assert server.started
        origin = f"http://127.0.0.1:{port}"
        with httpx.Client(base_url=origin, trust_env=False) as client:
            login = client.post("/api/session", json={"key": app.state.launch_key},
                                headers={"Origin": origin})
            assert login.status_code == 200
            cookie = client.cookies.get("sozvon_session")
        with socket.create_connection(("127.0.0.1", port), timeout=5) as connection:
            connection.sendall((f"GET /api/meetings/{mid}/export?format=pdf HTTP/1.1\r\n"
                                f"Host: 127.0.0.1:{port}\r\n"
                                f"Cookie: sozvon_session={cookie}\r\n\r\n").encode())
            assert rendering.wait(5)
            connection.shutdown(socket.SHUT_RDWR)
        assert finished.wait(5), "Real TCP disconnect must reach the export cancellation event"
        assert children and children[0].poll() is not None
        deadline = time.monotonic() + 3
        while service._export_lock.locked() and time.monotonic() < deadline:
            time.sleep(.01)
        assert not service._export_lock.locked()
        assert list((tmp_path / "exports" / "tmp").iterdir()) == []
        assert not service._export_stop.is_set()
    finally:
        service._export_stop.set()
        server.should_exit = True
        thread.join(5)
        listener.close()
        assert not thread.is_alive()


@pytest.mark.parametrize("fmt", ["md", "docx", "pdf"])
def test_stale_report_409_current_transcript_200_json_keeps_history(tmp_path, fmt):
    from pypdf import PdfReader

    app = app_with_router(tmp_path)
    repo = app.state.service.repo
    mid = repo.create_text("Синтетическая встреча", "Цитата старой версии")
    old = repo.detail(mid)
    sid = old["segments"][0]["id"]
    repo.save_report(mid, {"document": {"summary": [
        {"text": "Старая сводка", "evidence": [{"segment_id": sid, "quote": "Цитата старой версии"}]}]}})
    rid = repo.detail(mid)["reports"][0]["id"]
    repo.save_segments(mid, [{"text": "Текущая версия"}])
    before = repo.detail(mid)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        response = client.get(f"/api/meetings/{mid}/export?format={fmt}&report_id={rid}")
        assert response.status_code == 409 and "устарел" in response.json()["detail"]
        assert "content-disposition" not in response.headers
        transcript = client.get(f"/api/meetings/{mid}/export?format={fmt}&content=transcript")
        assert transcript.status_code == 200
        if fmt == "pdf":
            text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(transcript.content)).pages)
        elif fmt == "docx":
            text = "\n".join(p.text for p in Document(BytesIO(transcript.content)).paragraphs)
        else:
            text = transcript.text
        assert "Текущая версия" in text and "Старая сводка" not in text
        archive = client.get(f"/api/meetings/{mid}/export?format=json&include_quotes=0")
        assert archive.status_code == 200
        assert archive.json() == before
        assert repo.detail(mid) == before


def test_unknown_report_and_meeting_404_busy_409(tmp_path):
    app = app_with_router(tmp_path)
    service = app.state.service
    mid = service.repo.create_text("Ёж", "Текст")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        assert client.get(f"/api/meetings/{'f' * 32}/export").status_code == 404
        assert client.get(f"/api/meetings/{mid}/export?content=report").status_code == 404
        service._export_lock.acquire()
        try:
            response = client.get(f"/api/meetings/{mid}/export?format=pdf")
            assert response.status_code == 409
            assert "content-disposition" not in response.headers
        finally:
            service._export_lock.release()


@pytest.mark.parametrize("fmt", ["md", "docx", "pdf"])
def test_quote_opt_out_keeps_references_in_real_report_download(tmp_path, fmt):
    from pypdf import PdfReader

    app = app_with_router(tmp_path)
    service = app.state.service
    mid = service.repo.create_text("Русское имя Ёж", "Полный секретный текст цитаты")
    sid = service.repo.detail(mid)["segments"][0]["id"]
    service.repo.save_report(mid, {"document": {"tasks": [{"text": "Проверить описание",
        "owner": None, "due": None, "evidence": [
            {"segment_id": sid, "quote": "Полный секретный текст цитаты"}]}]}})
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        response = client.get(f"/api/meetings/{mid}/export?format={fmt}&include_quotes=0")
        assert response.status_code == 200
        if fmt == "pdf":
            text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages)
        elif fmt == "docx":
            text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        else:
            text = response.text
        assert "[Абзац 1] Источник" in text
        assert "Полный секретный текст цитаты" not in text
        assert "Проверить описание" in text and "Не назначен" in text


def test_report_missing_on_empty_meeting_is_404_not_validation_error(tmp_path):
    app = app_with_router(tmp_path)
    mid = app.state.service.repo.create("Без расшифровки", "audio")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        assert client.get(f"/api/meetings/{mid}/export?content=report").status_code == 404
        assert client.get(f"/api/meetings/{mid}/export?content=transcript").status_code == 422


def test_explicit_report_id_never_silently_falls_back_to_transcript(tmp_path):
    app = app_with_router(tmp_path)
    mid = app.state.service.repo.create_text("Ёж", "Текст")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        login(app, client)
        response = client.get(f"/api/meetings/{mid}/export?report_id={'f' * 32}")
        assert response.status_code == 404
