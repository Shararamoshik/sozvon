"""Independent readers verify synthetic DOCX/PDF exports."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from sozvon.export.document import Block, ExportDocument

ROOT = Path(__file__).resolve().parents[1]
FONTS = ROOT / "sozvon/web/static/fonts"


@pytest.fixture
def sample():
    return ExportDocument("Встреча — проект Ёж", (
        Block("h2", "Решения"),
        Block("p", "Анна & Илья: проверка <условия> выполнена."),
        Block("h2", "Задачи"),
        Block("p", "Обновить инструкцию"),
        Block("meta", "Исполнитель: Не назначен · Срок: Не указан"),
        Block("quote", "[Абзац 2] Срок пока не обещаю."),
    ))


def test_docx_is_real_document_with_russian_text(tmp_path, sample):
    from sozvon.export.docx import render_docx

    path = tmp_path / "result.docx"
    render_docx(sample, path)
    doc = Document(path)
    assert [p.text for p in doc.paragraphs] == [sample.title, *(b.text for b in sample.blocks)]
    assert doc.core_properties.title == sample.title
    assert doc.core_properties.author == "Созвон"
    section = doc.sections[0]
    assert abs(section.page_width - Cm(21)) < 1000
    assert abs(section.page_height - Cm(29.7)) < 1000
    assert abs(section.left_margin - Cm(2)) < 1000
    assert doc.paragraphs[1].style.name == "Heading 1"
    assert doc.paragraphs[1].paragraph_format.keep_with_next
    # OOXML stores paragraph indents in integer twips, unlike input EMUs.
    assert abs(doc.paragraphs[-1].paragraph_format.left_indent - Cm(.5)) < 1000
    assert doc.paragraphs[-2].runs[0].font.size.pt == 9
    assert doc.styles["Normal"].font.size.pt == 11
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        fonts = doc.styles[name].element.rPr.rFonts
        assert all(fonts.get(qn("w:" + key)) == "Arial"
                   for key in ("ascii", "hAnsi", "eastAsia", "cs"))
    assert section.footer._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')


def test_pdf_has_extractable_russian_text_and_embedded_fonts(tmp_path, sample):
    from pypdf import PdfReader

    from sozvon.export.pdf import render_pdf

    path = tmp_path / "result.pdf"
    render_pdf(sample, path, FONTS)
    assert path.read_bytes().startswith(b"%PDF-")
    reader = PdfReader(path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for expected in (sample.title, *(b.text for b in sample.blocks)):
        assert expected in text
    assert reader.metadata.title == sample.title
    assert reader.metadata.author == "Созвон"
    page = reader.pages[0]
    assert float(page.mediabox.width) == pytest.approx(595.28, abs=.1)
    assert float(page.mediabox.height) == pytest.approx(841.89, abs=.1)
    fonts = [font.get_object() for font in page["/Resources"]["/Font"].values()]
    embedded = [font for font in fonts if "/FontDescriptor" in font]
    assert len(embedded) == 2
    assert all(font["/FontDescriptor"]["/FontFile2"].get_data() for font in embedded)
    assert all(font["/ToUnicode"].get_data() for font in embedded)
    assert text.strip().startswith("1")  # Footer page number is selectable text too.


def test_missing_font_is_named_error_even_after_previous_render(tmp_path, sample):
    from sozvon.export.pdf import render_pdf

    render_pdf(sample, tmp_path / "first.pdf", FONTS)
    with pytest.raises(ValueError, match="GolosText-Regular.ttf"):
        render_pdf(sample, tmp_path / "missing.pdf", tmp_path)
    assert not (tmp_path / "missing.pdf").exists()


def test_pdf_many_pages_keep_last_paragraph(tmp_path):
    from pypdf import PdfReader

    from sozvon.export.pdf import render_pdf

    blocks = tuple(Block("p", f"Пункт {index}. " + "Длинный синтетический абзац. " * 35)
                   for index in range(110))
    sample = ExportDocument("Синтетическая многостраничная встреча", blocks + (
        Block("p", "ПОСЛЕДНИЙ АБЗАЦ — проверка завершена."),))
    path = tmp_path / "many.pdf"
    render_pdf(sample, path, FONTS)
    reader = PdfReader(path)
    assert len(reader.pages) >= 10
    text = "\n".join(page.extract_text() for page in reader.pages)
    for index in range(110):
        assert f"Пункт {index}." in text
    assert "ПОСЛЕДНИЙ АБЗАЦ — проверка завершена." in reader.pages[-1].extract_text()
    assert all(page.extract_text().strip().startswith(str(index))
               for index, page in enumerate(reader.pages, 1))


def test_very_long_word_and_single_paragraph_wrap_without_loss(tmp_path):
    from pypdf import PdfReader

    from sozvon.export.pdf import render_pdf

    word = "Длинноеслово" * 2000
    paragraph = "Большой единый абзац с переносами. " * 2000 + "КОНЕЦАБЗАЦА"
    path = tmp_path / "wrapped.pdf"
    render_pdf(ExportDocument("Перенос", (Block("p", word), Block("quote", paragraph))),
               path, FONTS)
    reader = PdfReader(path)
    # Remove each independently verified footer; page numbers can split a word.
    content = []
    for number, page in enumerate(reader.pages, 1):
        footer, body = page.extract_text().split("\n", 1)
        assert footer == str(number)
        content.append(body)
    text = "".join(content)
    compact = "".join(text.split())
    assert compact.count("Длинноеслово") == 2000
    assert compact.count("Большойединыйабзацспереносами.") == 2000
    assert "КОНЕЦАБЗАЦА" in text
    assert len(reader.pages) > 1


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_xml_image_markup_stays_literal_without_network(tmp_path, monkeypatch, fmt):
    import socket

    from pypdf import PdfReader

    from sozvon.export.docx import render_docx
    from sozvon.export.pdf import render_pdf

    def denied(*args, **kwargs):
        raise AssertionError("Renderer attempted network access")

    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket.socket, "connect", denied)
    literal = '<img src="https://example.invalid/test.png"/><script>test</script> & Ё'
    sample = ExportDocument("Текст <b>не разметка</b>", tuple(
        Block(kind, literal) for kind in ("h2", "p", "quote", "meta")))
    path = tmp_path / ("literal." + fmt)
    if fmt == "docx":
        render_docx(sample, path)
        text = "\n".join(p.text for p in Document(path).paragraphs)
    else:
        render_pdf(sample, path, FONTS)
        text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    compact = "".join(text.split())
    assert "".join(sample.title.split()) in compact
    assert compact.count("".join(literal.split())) == 4


def test_verification_cli_produces_readable_exports_and_page_images(tmp_path):
    import json
    import subprocess
    import sys

    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/verify_exports.py"), "--output", str(tmp_path)],
        cwd=ROOT, capture_output=True, text=True, timeout=60, check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["synthetic"] is True
    assert manifest["docx_visual_review"] == "not_performed"
    assert {item["filename"] for item in manifest["files"]} == {
        "short.docx", "short.pdf", "long.docx", "long.pdf"}
    for item in manifest["files"]:
        artifact = tmp_path / item["filename"]
        assert artifact.stat().st_size == item["bytes"] > 1000
        assert item["missing_strings"] == 0
        if artifact.suffix == ".pdf":
            assert item["pages"] >= 1
            assert len(item["images"]) == item["pages"]
            for name in item["images"]:
                assert (tmp_path / name).read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
        else:
            # DOCX pagination belongs to Word/LibreOffice, not python-docx.
            assert item["pages"] is None
