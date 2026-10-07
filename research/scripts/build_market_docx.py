"""Сборка DOCX из research/MARKET.md + приложение с диаграммами."""
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / 'MARKET.md'
OUT = ROOT / 'Исследование_рынка_Созвон.docx'
CHARTS = ['market_segmentation.png', 'adoption.png', 'market_growth.png',
          'competitors_map.png', 'cost.png', 'data_flow.png', 'pilot_plan.png']
NAVY = RGBColor(0x0F, 0x1B, 0x32)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)


def add_runs(paragraph, text):
    for token in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
        else:
            paragraph.add_run(token)


lines = MD.read_text(encoding='utf-8').splitlines()
index = 0
while index < len(lines):
    line = lines[index]
    stripped = line.strip()
    if not stripped:
        index += 1
        continue
    if stripped.startswith('| '):
        rows = []
        while index < len(lines) and lines[index].strip().startswith('|'):
            cells = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', cell) for cell in cells):
                rows.append(cells)
            index += 1
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = 'Light Grid Accent 1'
        for row_index, row in enumerate(rows):
            for col_index, cell in enumerate(row):
                paragraph = table.cell(row_index, col_index).paragraphs[0]
                add_runs(paragraph, cell)
                if row_index == 0:
                    for run in paragraph.runs:
                        run.bold = True
        doc.add_paragraph()
        continue
    if stripped == '---':
        index += 1
        continue
    if stripped.startswith('# '):
        header = doc.add_paragraph()
        run = header.add_run(stripped[2:])
        run.bold = True
        run.font.size = Pt(19)
        run.font.color.rgb = NAVY
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif stripped.startswith('## '):
        doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith('### '):
        doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith('- '):
        add_runs(doc.add_paragraph(style='List Bullet'), stripped[2:])
    elif re.match(r'^\d+\. ', stripped):
        add_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', stripped))
    else:
        add_runs(doc.add_paragraph(), stripped)
    index += 1

doc.add_page_break()
doc.add_heading('Приложение: диаграммы', level=1)
for name in CHARTS:
    path = ROOT / 'charts' / name
    doc.add_picture(str(path), width=Cm(16.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph()
    run = caption.add_run(name)
    run.italic = True
    run.font.size = Pt(9)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save(OUT)
print('DOCX_OK', OUT.name, OUT.stat().st_size, 'bytes, chart-images', len(CHARTS))
